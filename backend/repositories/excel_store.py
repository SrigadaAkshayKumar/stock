"""Excel-workbook backed repository used during development.

Each table in :data:`repositories.schema.TABLES` is a worksheet whose first
row holds the column names. Writes are serialised with a lock and the
workbook is saved after every mutation, which is fine for development
volumes. Production should use PostgreSQL (``db/schema.sql``).
"""

import os
import threading
import uuid
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from openpyxl import Workbook, load_workbook

from repositories.schema import TABLES


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


class ExcelStore:
    def __init__(self, path: str):
        self.path = path
        self._lock = threading.RLock()
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self._ensure_workbook()

    # -- internal -----------------------------------------------------------
    def _ensure_workbook(self) -> None:
        with self._lock:
            if os.path.exists(self.path):
                wb = load_workbook(self.path)
            else:
                wb = Workbook()
                wb.remove(wb.active)
            changed = False
            for table, columns in TABLES.items():
                if table not in wb.sheetnames:
                    wb.create_sheet(table).append(columns)
                    changed = True
            if changed or not os.path.exists(self.path):
                wb.save(self.path)

    def _rows(self, ws) -> List[Dict[str, Any]]:
        values = list(ws.values)
        if not values:
            return []
        header = list(values[0])
        return [dict(zip(header, row)) for row in values[1:]]

    @staticmethod
    def _matches(row: Dict[str, Any], filters: Dict[str, Any]) -> bool:
        return all(row.get(k) == v for k, v in filters.items())

    # -- public API ---------------------------------------------------------
    def insert(self, table: str, record: Dict[str, Any]) -> Dict[str, Any]:
        columns = TABLES[table]
        row = {c: record.get(c) for c in columns}
        row["id"] = row.get("id") or uuid.uuid4().hex
        if "created_at" in columns and not row.get("created_at"):
            row["created_at"] = utc_now()
        with self._lock:
            wb = load_workbook(self.path)
            wb[table].append([row[c] for c in columns])
            wb.save(self.path)
        return row

    def find(self, table: str, limit: Optional[int] = None, **filters) -> List[Dict[str, Any]]:
        with self._lock:
            wb = load_workbook(self.path, read_only=True)
            rows = [r for r in self._rows(wb[table]) if self._matches(r, filters)]
            wb.close()
        return rows[-limit:] if limit else rows

    def find_one(self, table: str, **filters) -> Optional[Dict[str, Any]]:
        rows = self.find(table, **filters)
        return rows[0] if rows else None

    def update(self, table: str, record_id: str, changes: Dict[str, Any]) -> bool:
        columns = TABLES[table]
        with self._lock:
            wb = load_workbook(self.path)
            ws = wb[table]
            for row in ws.iter_rows(min_row=2):
                if row[0].value == record_id:
                    for key, value in changes.items():
                        if key in columns:
                            row[columns.index(key)].value = value
                    wb.save(self.path)
                    return True
        return False
